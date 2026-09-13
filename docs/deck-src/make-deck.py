"""Build the llama-receipts deck: what exists, how it fits, where it goes.

    uv run --with python-pptx --with pillow python3 make-deck.py     # writes ../llama-receipts-activation-trace-secure-v2.pptx
    uv run --with pillow python3 render-shots.py                    # re-renders shots/*.png from shots/*.txt
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from pathlib import Path
OUT = str(Path(__file__).resolve().parent.parent / "llama-receipts-activation-trace-secure-v2.pptx")

INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x5F, 0x67, 0x70)
ACCENT = RGBColor(0x0B, 0x6E, 0x99)
ACCENT_SOFT = RGBColor(0xDC, 0xEC, 0xF4)
GREEN_SOFT = RGBColor(0xE2, 0xF3, 0xE4)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER_SOFT = RGBColor(0xFB, 0xEF, 0xD9)
AMBER = RGBColor(0x9A, 0x6A, 0x00)
RULE = RGBColor(0xD9, 0xDE, 0xE3)
OK = RGBColor(0x2E, 0x7D, 0x32)
BAD = RGBColor(0xB0, 0x2A, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF4, 0xF6, 0xF8)
FONT = "Helvetica Neue"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
M = Inches(0.7)


def txt(shape, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font=FONT):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return tf


def add_par(tf, text, size=18, bold=False, color=INK, level=0, space_before=6, font=FONT, align=None):
    p = tf.add_paragraph()
    p.level = level
    p.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return p


def header(slide, title, kicker=None):
    if kicker:
        k = slide.shapes.add_textbox(M, Inches(0.38), W - 2 * M, Inches(0.35))
        txt(k, kicker.upper(), 11, True, ACCENT)
    t = slide.shapes.add_textbox(M, Inches(0.7), W - 2 * M, Inches(1.0))
    txt(t, title, 28, True, INK)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, M, Inches(1.72), W - 2 * M, Emu(12700))
    line.fill.solid()
    line.fill.fore_color.rgb = RULE
    line.line.fill.background()


def footer(slide, n):
    f = slide.shapes.add_textbox(M, H - Inches(0.5), W - 2 * M, Inches(0.3))
    tf = txt(f, "llama-receipts, Sep 13, 2026", 10, False, MUTED)
    r = tf.paragraphs[0].add_run()
    r.text = f"     {n}"
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED
    r.font.name = FONT


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def bullets(slide, items, top=Inches(2.0), left=M, width=None, size=18, height=None):
    width = width or (W - 2 * M)
    height = height or (H - top - Inches(0.8))
    box_ = slide.shapes.add_textbox(left, top, width, height)
    tf = box_.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        level = 0
        if isinstance(it, tuple):
            it, level = it
        if first:
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = it
            r.font.size = Pt(size if level == 0 else size - 3)
            r.font.color.rgb = INK if level == 0 else MUTED
            r.font.name = FONT
            first = False
        else:
            add_par(tf, it, size if level == 0 else size - 3, color=INK if level == 0 else MUTED, level=level, space_before=10 if level == 0 else 3)
    return box_


def table(slide, rows, top, left=M, width=None, col_widths=None, size=12, header_fill=ACCENT_SOFT, row_h=Inches(0.36)):
    width = width or (W - 2 * M)
    nrows, ncols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrows, ncols, left, top, width, row_h * nrows)
    tbl = shp.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = int(width * cw / total)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            rr = p.add_run()
            rr.text = str(val)
            rr.font.size = Pt(size)
            rr.font.name = FONT
            rr.font.bold = r == 0
            color = INK
            if r > 0 and val in ("accept", "reject", "proved", "checked"):
                color = OK if val in ("accept", "proved", "checked") else BAD
                rr.font.bold = True
            rr.font.color.rgb = color
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else WHITE
    return tbl


def box(slide, x, y, w, h, text, size=13, fill=WHITE, line=ACCENT, bold=False, color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return s


def sub(shape, text, size=11, color=MUTED, align=PP_ALIGN.CENTER):
    add_par(shape.text_frame, text, size, color=color, space_before=3, align=align)


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


SH = str(Path(__file__).resolve().parent / "shots")


def fit_pic(slide, name, top_in, max_h_in, max_w_in=11.9):
    """Place a screenshot scaled to fit the box, centred horizontally."""
    from PIL import Image
    path = f"{SH}/{name}.png"
    w, h = Image.open(path).size
    aspect = h / w
    width_in = min(max_w_in, max_h_in / aspect)
    left = (13.333 - width_in) / 2
    return slide.shapes.add_picture(path, Inches(left), Inches(top_in), width=Inches(width_in))


n = 0


def new_slide(title, kicker=None):
    global n
    n += 1
    s = prs.slides.add_slide(BLANK)
    if title:
        header(s, title, kicker)
    footer(s, n)
    return s


# 1. title ------------------------------------------------------------------------------
s = new_slide(None)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
band.fill.solid()
band.fill.fore_color.rgb = INK
band.line.fill.background()
t = s.shapes.add_textbox(M, Inches(2.1), W - 2 * M, Inches(1.4))
txt(t, "Inference receipts for llama.cpp", 44, True, WHITE)
t2 = s.shapes.add_textbox(M, Inches(3.4), W - 2 * M, Inches(1.5))
tf = txt(t2, "Signed receipts, a committed trace of every tensor, a verifier that never runs the model,", 20, False, RGBColor(0xC9, 0xD6, 0xE0))
add_par(tf, "a formal specification in Lean, and a first integrity proof of the native arithmetic", 20, color=RGBColor(0xC9, 0xD6, 0xE0), space_before=2)
t3 = s.shapes.add_textbox(M, Inches(5.7), W - 2 * M, Inches(0.9))
tf = txt(t3, "Luis Cosio, SL5. Sep 13, 2026", 14, False, RGBColor(0xC9, 0xD6, 0xE0))
add_par(tf, "github.com/luiscosio/llama.cpp, pull requests 1, 2 and 3", 14, color=RGBColor(0xC9, 0xD6, 0xE0), space_before=4)
notes(s, "What exists as of Sep 13, 2026, how the pieces fit, and the path to zero-knowledge proofs of llama.cpp inference.")

# 2. problem ----------------------------------------------------------------------------
s = new_slide("You get tokens back and nothing else", "The problem")
bullets(s, [
    "When someone else runs a model for you, you cannot tell which model ran, at which quantization, on which prompt.",
    "Re-running everything yourself works, but costs as much as the original run.",
    "A zero-knowledge proof would settle it. None exists today for llama.cpp's native arithmetic.",
    "So we built the ladder that leads there: a record of exactly what ran, cheap to check, and the pieces a proof system plugs into.",
], size=19)
notes(s, "Framing for people outside the field. Everything after this slide is what exists in code.")

# 3. system diagram ---------------------------------------------------------------------
s = new_slide("What we built, in one picture", "Overview")
y = Inches(2.15)
bw, bh = Inches(2.55), Inches(1.15)
gap = Inches(0.45)
x0 = M
b1 = box(s, x0, y, bw, bh, "llama.cpp prover", 15, ACCENT_SOFT, ACCENT, True)
sub(b1, "generates, hashes every tensor, writes receipt and trace")
x1 = x0 + bw + gap
b2 = box(s, x1, y, bw, bh, "Receipt + trace", 15, WHITE, ACCENT, True)
sub(b2, "tokens, sampler, model hashes, Merkle root, openings")
x2 = x1 + bw + gap
b3 = box(s, x2, y, bw, bh, "Trace verifier", 15, ACCENT_SOFT, ACCENT, True)
sub(b3, "Python, numpy, no model run; pinned graph topology")
x3 = x2 + bw + gap
b4 = box(s, x3, y, bw, bh, "Replay verifier", 15, ACCENT_SOFT, ACCENT, True)
sub(b4, "re-runs the model, seeded sampler, two-regime rule")
arrow(s, x0 + bw, y + bh // 2, x1, y + bh // 2)
arrow(s, x1 + bw, y + bh // 2, x2, y + bh // 2)
arrow(s, x2 + bw, y + bh // 2, x3, y + bh // 2)
# second row: spec and proof
y2 = y + bh + Inches(1.35)
b5 = box(s, x1, y2, bw, bh, "Lean specification", 15, GREEN_SOFT, GREEN, True)
sub(b5, "formats, arithmetic, protocol; theorems; checker")
b6 = box(s, x2, y2, bw, bh, "GKR proof of a node", 15, AMBER_SOFT, AMBER, True)
sub(b6, "Expander, M31 field, weights bound to GGUF")
arrow(s, x1 + bw // 2, y + bh, x1 + bw // 2, y2)
arrow(s, x2 + bw // 2, y + bh, x2 + bw // 2, y2)
cap = s.shapes.add_textbox(M, y2 + bh + Inches(0.25), W - 2 * M, Inches(0.8))
tf = txt(cap, "The receipt is the contract. The trace verifier checks it without the model. The Lean spec says what the receipt means and checks the same artifact independently. The proof replaces re-execution of a node with a cryptographic check of its arithmetic.", 13, False, MUTED)
notes(s, "Four programs and one document. Prover and replay verifier are C++ in llama.cpp; the trace verifier is Python; the spec is Lean 4; the proof is Rust on Expander.")

# 4. receipt ----------------------------------------------------------------------------
s = new_slide("A receipt records exactly what ran", "Receipt")
bullets(s, [
    "Model: SHA-256 of the GGUF file and of every tensor, plus a Merkle root over the tensors.",
    "Request: prompt text and tokens, temperature, top-k, top-p, seed.",
    "Response: tokens, and per position the log-prob, the uniform draw, its distance to the decision boundary, and a hash of the logits.",
    "Commitment: SHA-256 over prompt and response tokens.",
    "The sampler is replayable by construction:",
], width=Inches(7.2), size=17)
b = box(s, Inches(8.3), Inches(2.3), Inches(4.3), Inches(1.1), "u(t) = SHA-256(\"llama-receipts/u/v1/\" || seed || t)", 15, ACCENT_SOFT, ACCENT, True, ACCENT)
b2 = s.shapes.add_textbox(Inches(8.3), Inches(3.55), Inches(4.3), Inches(2.6))
txt(b2, "Ties break by token id. Top-k and top-p run in double precision. Anyone can reproduce the draw without RNG state, so a receipt shows how far each choice sat from the boundary, not only which token won.", 14, False, MUTED)
notes(s, "Signing is left to the operator's key tooling; the README shows an openssl one-liner.")

# 5. trace diagram ----------------------------------------------------------------------
s = new_slide("Every computing node becomes a leaf that names its producers", "Trace")
y0 = Inches(2.2)
nodes = ["token embedding", "attention block", "feed-forward block", "logits"]
xs = M
bw, bh = Inches(2.3), Inches(0.62)
for i, name in enumerate(nodes):
    box(s, xs, y0 + i * Inches(0.95), bw, bh, name, 13)
    if i < len(nodes) - 1:
        arrow(s, xs + bw // 2, y0 + i * Inches(0.95) + bh, xs + bw // 2, y0 + (i + 1) * Inches(0.95))
cap = s.shapes.add_textbox(xs, y0 + 4 * Inches(0.95) - Inches(0.2), bw, Inches(0.6))
txt(cap, "GGML graph, observed through cb_eval", 11, False, MUTED, PP_ALIGN.CENTER)
arrow(s, xs + bw, y0 + Inches(1.6), Inches(3.55), y0 + Inches(1.6))
lx, lw = Inches(3.65), Inches(4.45)
card = box(s, lx, y0, lw, Inches(3.7), "", 12, WHITE, RULE, shape=MSO_SHAPE.RECTANGLE, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)
tf = card.text_frame
p = tf.paragraphs[0]
p.runs[0].text = "leaf"
p.runs[0].font.bold = True
p.runs[0].font.color.rgb = ACCENT
for line in [
    "graph g, node i, name, op, op params",
    "out: type, shape, strides, offset,",
    "     hash of the base tensor",
    "srcs: weight: name, type, shape, GGUF hash",
    "      data: shape, offset, base hash,",
    "            producer = leaf that wrote it",
    "            (-1 graph input, -2 zeroed KV cache)",
]:
    q = add_par(tf, line, 10.5, color=INK, space_before=3, font="Menlo")
    q.alignment = PP_ALIGN.LEFT
cap2 = s.shapes.add_textbox(lx, y0 + Inches(3.8), lw, Inches(0.5))
txt(cap2, "Layout ops (reshape, view, permute) are resolved to what they view. Inputs are hashed over their whole base tensor, so the edge check is a string comparison.", 11, False, MUTED)
arrow(s, lx + lw, y0 + Inches(1.6), Inches(8.45), y0 + Inches(1.6))
tx = Inches(8.5)
leaf_w, leaf_h = Inches(0.9), Inches(0.42)
gap = Inches(0.18)
leaf_y = y0 + Inches(2.6)
leaf_xs = [tx + i * (leaf_w + gap) for i in range(4)]
for i, lxx in enumerate(leaf_xs):
    box(s, lxx, leaf_y, leaf_w, leaf_h, f"leaf {i}", 10, ACCENT_SOFT, ACCENT)
mid_y = leaf_y - Inches(0.95)
mids = [(leaf_xs[0] + leaf_xs[1] + leaf_w) // 2 - leaf_w // 2, (leaf_xs[2] + leaf_xs[3] + leaf_w) // 2 - leaf_w // 2]
for mx in mids:
    box(s, mx, mid_y, leaf_w, leaf_h, "h(0x01||l||r)", 9, WHITE, ACCENT)
for i, lxx in enumerate(leaf_xs):
    arrow(s, lxx + leaf_w // 2, leaf_y, mids[i // 2] + leaf_w // 2, mid_y + leaf_h)
root_y = mid_y - Inches(0.95)
root_x = (mids[0] + mids[1]) // 2
box(s, root_x, root_y, leaf_w, leaf_h, "root", 11, ACCENT, ACCENT, True, WHITE)
for mx in mids:
    arrow(s, mx + leaf_w // 2, mid_y, root_x + leaf_w // 2, root_y + leaf_h)
cap3 = s.shapes.add_textbox(tx - Inches(0.1), leaf_y + Inches(0.55), Inches(4.4), Inches(1.5))
txt(cap3, "The root goes into the receipt. A second pass replays the same tokens, reproduces the root bit for bit, and captures the bytes of a Fiat-Shamir sample of nodes: the openings.", 11, False, MUTED)
notes(s, "23,000 leaves for a 32-token generation of a 1.5B model. Same-backend replay is deterministic once the KV cache starts zeroed.")

# 6. verifier checks ----------------------------------------------------------------------
s = new_slide("Trace verifier checks without model execution", "Verifier")
rows = [
    ["Check", "What it establishes", "Needs tensor bytes"],
    ["1. Root + content", "Leaves, token arrays and displayed text match their commitments", "no"],
    ["2. Topology + challenge", "Graph structure matches verifier policy; openings are unpredictable and meet its minimum", "no"],
    ["3. Structure", "One graph per token, prefill included", "no"],
    ["4. Inputs", "Token, position, row-selection, KV-cell and mask inputs hash to what the receipt claims", "no"],
    ["5. Edges", "Every input equals its producer's output; the KV cache started as zeros", "no"],
    ["6. Logits binding", "Each graph's logits hash equals the receipt's per-token logits hash", "no"],
    ["7. Openings", "Path to root, bytes hash to the leaf, weights match the verifier's GGUF, numpy re-execution within tolerance", "yes, for k nodes"],
]
table(s, rows, Inches(2.05), col_widths=[1.7, 6.2, 1.6], size=13, row_h=Inches(0.55))
notes(s, "Checks 1 to 6 are hash comparisons over about 23,000 small JSON objects. Only the openings touch tensor bytes, and there are k of them, 32 by default.")

# 7. arithmetic references ----------------------------------------------------------------
s = new_slide("Re-execution matches the kernel that ran", "Verifier")
bullets(s, [
    "Matmuls are judged against three references and the closest one counts:",
    ("f32: dequantized weight times float32 activation. Metal decode kernels.", 1),
    ("f16: both operands rounded to half. Metal prefill kernels.", 1),
    ("q8k: ggml's CPU path. Activation quantized to Q8_K, then an integer dot product with the 4-bit or 6-bit weight blocks and per-sub-block scales. Reproduced exactly.", 1),
    "The q8k path is the statement a proof system constrains, and it is what the Lean spec and the GKR circuit use.",
], width=Inches(6.6), size=16)
rows = [
    ["Op", "Metal", "CPU", "Limit"],
    ["ADD, MUL, RMS_NORM, SWIGLU", "1e-7", "1e-7", "1e-5"],
    ["ROPE", "4e-7", "7e-7", "5e-5"],
    ["GET_ROWS, CONT, SET_ROWS", "exact", "exact", "exact"],
    ["MUL_MAT, decode", "4e-7 (f32)", "8e-7 (q8k)", "8e-3"],
    ["MUL_MAT, prefill", "2.2e-3 (f16)", "1e-7 (q8k)", "8e-3"],
    ["FLASH_ATTN_EXT", "1.9e-3", "7.5e-3", "2e-2"],
]
table(s, rows, Inches(2.05), left=Inches(7.6), width=Inches(5.05), col_widths=[2.6, 1.3, 1.3, 1.0], size=11, row_h=Inches(0.42))
cap = s.shapes.add_textbox(Inches(7.6), Inches(5.1), Inches(5.05), Inches(1.0))
txt(cap, "Normalized error, max|out - ref| / max|ref|, qwen2.5 1.5B Q4_K_M on an Apple M5, several hundred openings per backend.", 11, False, MUTED)
notes(s, "Calibrated with the native llama.cpp tool on both backends.")

# 8. results ------------------------------------------------------------------------------
s = new_slide("Results with the native tool", "Results")
rows = [
    ["Case", "Verdict", "What decided it"],
    ["Honest, Metal, flash attention on", "accept", "all checks pass, worst matmul error 1.2e-3"],
    ["Honest, CPU, 8 threads", "accept", "matmuls match the integer reference to 5e-7"],
    ["Replay of the Metal receipt on Metal", "accept", "bit-exact, every logits hash matches"],
    ["Replay of the Metal receipt on CPU", "accept", "drift regime, token match 0.75, largest draw gap 0.08"],
    ["Q2_K model served, receipt claims Q4_K_M", "reject", "8 of 8 sampled weight matmuls fail, errors 0.2 to 0.4"],
    ["Honest receipt checked against the wrong file", "reject", "model file hash, then weight bindings"],
    ["One response token edited, commitments refreshed", "reject", "challenge indices differ, token input mismatches"],
    ["Interactive seed, wrong seed given to the verifier", "reject", "recorded seed is not the verifier's"],
]
table(s, rows, Inches(2.05), col_widths=[3.6, 0.9, 5.0], size=12, row_h=Inches(0.45))
cap = s.shapes.add_textbox(M, Inches(6.2), W - 2 * M, Inches(0.5))
txt(cap, "Python prototype on the same kernels, 4 prompts, 32 tokens, 32 openings: 28 of 28 verdicts correct.", 12, False, MUTED)
notes(s, "The Q2_K case forged the leaf metadata to claim Q4_K_M types; only the opened bytes betray it.")


s = new_slide('How it looks: proving with a trace, then replaying it', 'Screenshots')
p1 = fit_pic(s, "prove", 1.95, 2.75)
fit_pic(s, "replay", 1.95 + p1.height / 914400 + 0.25, 1.25)
cap = s.shapes.add_textbox(M, Inches(6.35), W - 2 * M, Inches(0.6))
txt(cap, 'Top: the prover generates, hashes every tensor, writes the receipt, replays once for the openings and writes the trace. Bottom: replay verification on the same backend is bit-exact.', 12, False, MUTED)

s = new_slide('How it looks: verifying', 'Screenshots')
fit_pic(s, "verify", 1.95, 4.3)
cap = s.shapes.add_textbox(M, Inches(6.35), W - 2 * M, Inches(0.6))
txt(cap, 'The trace verifier: no model execution, nineteen checks, every one passing, about four seconds on a laptop.', 12, False, MUTED)

# 9. Lean spec ---------------------------------------------------------------------------
s = new_slide("The specification in Lean 4: readable, provable, runnable", "Formal spec")
bullets(s, [
    "Specified, as executable definitions:",
    ("Byte-exact Q4_K, Q6_K and Q8_K block formats and the nibble order the kernels read.", 1),
    ("ggml's Q8_K activation quantization in Float32, step for step.", 1),
    ("The integer core s1, s2 and the Q6_K dot; the per-block float combination.", 1),
    ("Merkle tree, audit paths, Fiat-Shamir sample, canonical leaf bytes, edge and input checks. SHA-256 included.", 1),
    "Proved, without Mathlib or sorry:",
    ("Every decoded bit field is in range. Every s1 is below 2^25 and every s2 below 2,048,256, so the M31 circuit never wraps.", 1),
    ("Roots and audit paths for one- and two-leaf trees. Sampled indices are valid and never repeat.", 1),
], width=Inches(6.7), size=15)
rows = [
    ["spec-check on a real CPU receipt", "Result"],
    ["Merkle root over 8,086 leaves", "checked"],
    ["Pinned topology and content commitments", "checked"],
    ["Fiat-Shamir sample, 32 openings", "checked"],
    ["Data edges, 10,621", "checked"],
    ["Token and position inputs, 13 and 728", "checked"],
    ["Q8_K quants, scale bits, s1, s2, Q6_K dot", "checked"],
    ["Row value in Float vs float64 reference", "checked"],
]
table(s, rows, Inches(2.05), left=Inches(7.7), width=Inches(4.95), col_widths=[3.4, 1.2], size=12, row_h=Inches(0.42))
cap = s.shapes.add_textbox(Inches(7.7), Inches(5.15), Inches(4.95), Inches(1.0))
txt(cap, "The checker recomputes a real trace in under half a second and agrees with the C++ prover exactly. Stated as the next target: the (1 - f)^k sampling bound.", 11, False, MUTED)
notes(s, "Lean 4.33.1, lake build in seconds. Floating point is opaque to Lean's logic; the float parts are checked by differential testing.")


s = new_slide('How it looks: the specification in Lean', 'Screenshots')
fit_pic(s, "lean", 1.95, 4.3)
cap = s.shapes.add_textbox(M, Inches(6.35), W - 2 * M, Inches(0.6))
txt(cap, 'The integer sum s1 as Lean defines it, and the machine-checked theorem that every s1 stays below 2^25 in magnitude, which is what lets the proof circuit use Mersenne-31 field arithmetic as plain integers.', 12, False, MUTED)

s = new_slide('How it looks: the spec checked against a real trace', 'Screenshots')
fit_pic(s, "spec", 1.95, 4.3)
cap = s.shapes.add_textbox(M, Inches(6.35), W - 2 * M, Inches(0.6))
txt(cap, 'The Lean checker recomputes the root, the topology digest, the content commitments, the challenge, all edges and the token and position inputs of a real trace, then reproduces the Q8_K quants, scales, integer sums, row values and canonical JSON bytes on vectors from the GGUF and from Python.', 12, False, MUTED)

# 10. proof of one node ----------------------------------------------------------------
s = new_slide("Proofs of one node: integrity with Expander, zero knowledge with Groth16", "Proof")
y = Inches(2.1)
bw, bh = Inches(2.35), Inches(1.25)
gap = Inches(0.35)
x = M
c1 = box(s, x, y, bw, bh, "opened node", 14, WHITE, ACCENT, True)
sub(c1, "weight bytes from the GGUF, activation bytes from the trace")
x2 = x + bw + gap
c2 = box(s, x2, y, bw, bh, "witness", 14, WHITE, ACCENT, True)
sub(c2, "private: q4, scales, mins, under a registered commitment\npublic: q8, s1, s2")
x3 = x2 + bw + gap
c3 = box(s, x3, y, bw, bh, "GKR circuit, M31", 14, AMBER_SOFT, AMBER, True)
sub(c3, "s1 = Σ sc·Σ q4·q8\ns2 = Σ bsums·mn")
x4 = x3 + bw + gap
c4 = box(s, x4, y, bw, bh, "proof + commitment", 14, AMBER_SOFT, AMBER, True)
sub(c4, "the proof opens with the commitment to the private weights")
for a_, b_ in ((c1, c2), (c2, c3), (c3, c4)):
    arrow(s, a_.left + a_.width, y + bh // 2, b_.left, y + bh // 2)
y2 = y + bh + Inches(0.6)
v1 = box(s, x3, y2, bw, bh, "verify proof", 14, GREEN_SOFT, GREEN, True)
sub(v1, "holds no weights; compares the proof's commitment with the registered one")
v2 = box(s, x4, y2, bw, bh, "float step", 14, GREEN_SOFT, GREEN, True)
sub(v2, "Σ d_a·(d·s1 − dmin·s2), compare to the opened output")
arrow(s, x4 + bw // 2, y + bh, x4 + bw // 2, y2)
arrow(s, x4, y2 + bh // 2, x3 + bw, y2 + bh // 2)
rows = [
    ["Prover", "Node", "Private inputs", "Register", "Prove", "Proof", "Verify", "Zero knowledge"],
    ["Expander GKR, M31", "Vcur-22, 256 x 1536, whole node", "417,792", "2.0 s", "0.28 s", "6.5 MB", "0.04 s", "no: binding only"],
    ["Groth16, BN254", "Vcur-22, 16 groups of 16 rows", "26,112 per group", "Poseidon, 0.1 s", "3 s per group", "805 B per group", "0.16 s per group", "yes"],
]
table(s, rows, y2 + bh + Inches(0.35), col_widths=[1.5, 2.2, 1.2, 1.1, 0.9, 1.0, 1.0, 1.1], size=10, row_h=Inches(0.34))
notes(s, "One activation row from a CPU trace of qwen2.5 1.5B Q4_K_M, Apple M5, Sep 13, 2026. The registered commitment is Orion's root over the private input layer, computed from the GGUF alone. Not zero-knowledge yet: the commitment binds but does not hide, and the GKR has no masking. The Lean theorems s1_lt_2_pow_25 and s2_bound are what make M31 field arithmetic equal integer arithmetic here.")


s = new_slide('How it looks: proving one node', 'Screenshots')
fit_pic(s, "zk", 1.95, 4.3)
cap = s.shapes.add_textbox(M, Inches(6.35), W - 2 * M, Inches(0.6))
txt(cap, 'A 256 by 1536 value projection: the commitment registered from the GGUF, a proof carrying the same commitment, verification without weights in 0.04 seconds, three negative cases rejected, float step within 4e-7.', 12, False, MUTED)

# 11. where it lives -----------------------------------------------------------------------
s = new_slide("Where it lives: one fork, three stacked pull requests", "Code")
y = Inches(2.3)
bw, bh = Inches(2.7), Inches(0.9)
gap = Inches(0.5)
x = M
m0 = box(s, x, y, bw, bh, "ggml-org/llama.cpp master", 13, PANEL, RULE, True)
x1 = x + bw + gap
p1 = box(s, x1, y, bw, bh, "PR 1: receipts-trace", 13, ACCENT_SOFT, ACCENT, True)
sub(p1, "llama-receipts, verify_trace.py")
x2 = x1 + bw + gap
p2 = box(s, x2, y - Inches(0.75), bw, bh, "PR 2: receipts-lean-spec", 13, GREEN_SOFT, GREEN, True)
sub(p2, "spec/, Lean 4")
p3 = box(s, x2, y + Inches(0.75), bw, bh, "PR 3: receipts-zk-poc", 13, AMBER_SOFT, AMBER, True)
sub(p3, "zk/, Rust on Expander")
arrow(s, x + bw, y + bh // 2, x1, y + bh // 2)
arrow(s, x1 + bw, y + bh // 2, x2, y - Inches(0.75) + bh // 2)
arrow(s, x1 + bw, y + bh // 2, x2, y + Inches(0.75) + bh // 2)
cap = s.shapes.add_textbox(M, Inches(4.4), W - 2 * M, Inches(0.7))
txt(cap, "Fork luiscosio/llama.cpp, level with upstream master. All three are drafts inside the fork. Expander is used through a fork with two small macOS build fixes.", 13, False, MUTED)
code = box(s, M, Inches(5.1), W - 2 * M, Inches(1.5), "", 11, PANEL, RULE, shape=MSO_SHAPE.RECTANGLE, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)
tf = code.text_frame
p = tf.paragraphs[0]
p.runs[0].text = "llama-receipts -m model.gguf -p \"...\" -n 32 --seed 42 --trace --out receipt.json        # prove, trace"
p.runs[0].font.name = "Menlo"
for line in [
    "python3 verify_trace.py receipt.json --model model.gguf --expected-topology-sha256 HASH  # verify",
    "lake build && spec-check trace receipt.json receipt.trace.json HASH                      # Lean checker",
    "python3 zk_node.py receipt.json --model model.gguf                                       # prove one node",
]:
    q = add_par(tf, line, 11, color=INK, space_before=2, font="Menlo")
    q.alignment = PP_ALIGN.LEFT
notes(s, "Draft PRs in the fork so the work is reviewable as a unit. Upstream submission is a separate step with its own process.")

# 12. roadmap ----------------------------------------------------------------------------
s = new_slide("From one node to a proof of the whole forward pass", "Next")
y = Inches(2.15)
bw, bh = Inches(2.75), Inches(1.5)
gap = Inches(0.33)
steps = [
    ("Stage 1, done", "statement frozen, backend decided, budgets set, circuit checked against the Lean spec, weights bound to a registered commitment", GREEN_SOFT, GREEN),
    ("Stage 2, done", "registration manifest: digests, tensor table, tokenizer, execution spec, Orion and Poseidon commitments; Qwen3-0.6B registered", GREEN_SOFT, GREEN),
    ("Stage 3, done", "zero-knowledge proof of one node with Groth16: private weights under a Poseidon commitment, 805-byte proofs, verified in a browser", GREEN_SOFT, GREEN),
    ("Stages 4 to 6", "fixed-point mode: 20 of 20 tokens agree; isolation and adversarial tests run; registry and verify page live as a preview. Not built: the whole-token prover", AMBER_SOFT, AMBER),
]
for i, (k, d, fill, line) in enumerate(steps):
    x = M + i * (bw + gap)
    b = box(s, x, y, bw, bh, k, 16, fill, line, True)
    sub(b, d, 12, INK)
    if i < len(steps) - 1:
        arrow(s, x + bw, y + bh // 2, x + bw + gap, y + bh // 2)
bullets(s, [
    "Plan and records: PLAN.md and TODO.md, docs/stage1 to stage5. First model: Qwen3-0.6B requantized to Q4_K_M, every matmul a multiple of the 256-block.",
    "A whole token is 322 million weight nibbles: 1.6 billion R1CS constraints, or 320 million multiplications with lookups. Per-node Groth16 cannot scale to it; a GKR with lookups and masking can.",
    "Next: port the fixed-point mode into llama.cpp's kernels, add masking or aggregation on the GKR side, replace the proof-of-concept setup with ceremony parameters.",
], top=Inches(4.15), size=15)
notes(s, "Open provers today reach minutes per token at 7B; the closed claim is tens of tokens per second at 31B. The point here is that the native arithmetic is provable at all, which was not established before.")

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
